"""Purpose Statement export (.docx) — the readable Phase 0 artifact.

Two copies, one builder (ADR-010 as amended by ADR-018):

* OWNER copy (default): the statement plus Appendix A, the full verbatim
  interview transcript from the restricted store. For the program owner's
  own record; it is marked as internal.
* SHARE copy (share=True): the statement only. The transcript never leaves
  the restricted store in a share copy, and neither do verbatim answers.

Derived artifact only: governed/purpose_statement.json and
restricted/interview.json are the sources of truth, and this document can be
regenerated at any time.
"""
from __future__ import annotations

import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

from .defect_report import ACC, DIM, INK, _load, _para, _shade, styled_document

COMPLETE_MARKER = "[INTERVIEW-COMPLETE]"
WARN = RGBColor(0x9A, 0x5C, 0x00)
MARKUP = re.compile(r"\*\*(.+?)\*\*|(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])")


class NoPurposeStatement(Exception):
    pass


def _txt(x) -> str:
    """Best-effort text from the synthesis's mixed shapes (str | {text} | {claim} ...)."""
    if x is None:
        return ""
    if isinstance(x, str):
        return x.strip()
    if isinstance(x, dict):
        for k in ("text", "statement", "claim", "description", "note", "authority"):
            if isinstance(x.get(k), str) and x[k].strip():
                return x[k].strip()
        return ""
    return str(x)


def _date(ts) -> str:
    if not ts:
        return ""
    try:
        d = datetime.fromisoformat(str(ts).replace("Z", "+00:00"))
        return d.strftime("%-d %B %Y")
    except ValueError:
        return str(ts)[:10]


def _who(p) -> str:
    if not isinstance(p, dict):
        return _txt(p)
    name, cap = p.get("name") or "", p.get("capacity") or p.get("role") or ""
    return f"{name} ({cap})" if name and cap else (name or cap)


def _rich(paragraph, text: str, *, size=None, color=None, italic=False) -> None:
    """Add text to a paragraph, rendering the transcript's **bold** and *italic* markdown."""
    def run(t, b=False, i=False):
        r = paragraph.add_run(t); r.bold = b or None; r.italic = (italic or i) or None
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color
    pos = 0
    for m in MARKUP.finditer(text):
        if m.start() > pos:
            run(text[pos:m.start()])
        if m.group(1) is not None:
            run(m.group(1), b=True)
        else:
            run(m.group(2), i=True)
        pos = m.end()
    if pos < len(text):
        run(text[pos:])


def _bullets(doc, items: list[str]) -> None:
    for it in items:
        if it:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            _rich(p, it)


def _kv_table(doc, rows: list[tuple[str, str]]) -> None:
    rows = [(k, v) for k, v in rows if v]
    if not rows:
        return
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in rows:
        cells = t.add_row().cells
        cells[0].text = k; _shade(cells[0], "F3F4F6")
        cells[0].paragraphs[0].runs[0].bold = True
        cells[1].text = v
    t.autofit = False
    t.columns[0].width, t.columns[1].width = Cm(4.2), Cm(11.8)
    for r_ in t.rows:
        r_.cells[0].width, r_.cells[1].width = Cm(4.2), Cm(11.8)
        for c in r_.cells:
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    _para(doc, after=4)


def transcript(program_dir: Path) -> list[dict]:
    p = Path(program_dir) / "restricted" / "interview.json"
    if not p.exists():
        return []
    try:
        data = json.loads(p.read_text())
    except ValueError:
        return []
    return [m for m in data if isinstance(m, dict) and m.get("content")]


def build_docx(program_dir: str | Path, program_id: str, *, share: bool = False) -> bytes:
    pdir = Path(program_dir)
    g = pdir / "governed"
    ps = _load(g / "purpose_statement.json")
    if not ps:
        raise NoPurposeStatement("No Purpose Statement yet. Finish the interview and synthesize it first.")
    syn = ps.get("synthesis") or {}
    status = ps.get("status") or "draft"
    ratified = status == "ratified"
    rat = ps.get("ratification") or {}
    amends = ps.get("amendments") or []
    mode = ((syn.get("recommended_mode") or {}).get("mode") or "").strip()
    today = datetime.now(timezone.utc).strftime("%-d %B %Y")
    turns = [] if share else transcript(pdir)

    copy = "share copy" if share else "owner copy · internal"
    doc = styled_document(f"{program_id} — Purpose Statement v{ps.get('version', '')} · {copy} · page ", pdir)

    doc.add_heading(f"{program_id} — Purpose Statement", level=0)
    meta = [f"Version {ps.get('version', '')}",
            ("Ratified " + _date(rat.get("timestamp")) + " by " + _who(rat.get("ratified_by"))) if ratified
            else "Awaiting ratification"]
    if amends:
        a = amends[-1]
        meta.append(f"scope amended {_date(a.get('timestamp'))} ({a.get('amendment_id', '')})")
    meta.append(f"generated {today}")
    _para(doc, " · ".join(m for m in meta if m), size=9, color=DIM)

    # Copy banner: which copy this is and what it contains.
    if share:
        _para(doc, ("Share copy. This copy carries the synthesized statement only. The verbatim interview "
                    "transcript stays in the program owner's restricted record and is not included (ADR-010)."),
              size=9, color=DIM, italic=True, after=10)
    else:
        p = _para(doc, after=10)
        r = p.add_run("Owner copy — internal. "); r.bold = True; r.font.color.rgb = WARN; r.font.size = Pt(9)
        r = p.add_run(("Appendix A reproduces the full interview transcript (ADR-010 as amended by ADR-018). "
                       "Do not circulate this copy; use the share copy for anyone outside the program."))
        r.font.size = Pt(9); r.font.color.rgb = WARN

    doc.add_heading("1. Scope", level=1)
    _rich(_para(doc, after=8), _txt(syn.get("scope_sentence")) or "(scope not yet synthesized)")

    doc.add_heading("2. Decision served", level=1)
    _rich(_para(doc, after=8), _txt(syn.get("decision_served")) or "(not stated)")

    doc.add_heading("3. Mode", level=1)
    rm = syn.get("recommended_mode") or {}
    mode_line = {"redesign": "Redesign. The program resolves conflicts as binding decisions under its Principal's mandate.",
                 "refactor": "Refactor. The program codifies, clarifies and fills gaps; changes of substance are out of bounds."
                 }.get(mode, mode or "(not stated)")
    _para(doc, mode_line, after=4)
    if _txt(rm.get("mode_consistency_note")):
        _para(doc, _txt(rm.get("mode_consistency_note")), size=9, color=DIM, after=8)

    doc.add_heading("4. Who it serves", level=1)
    roles = ps.get("roles") or {}
    client = syn.get("client") or {}
    _kv_table(doc, [
        ("Principal", _who(roles.get("principal"))),
        ("Authority", _txt(roles.get("authority")) or _txt(client.get("authority"))),
        ("Program Owner", _who(roles.get("program_owner"))),
        ("Respondent", _who(roles.get("respondent"))),
        ("Consumers", "; ".join(_txt(c) for c in (syn.get("consumers") or []) if _txt(c))),
    ])

    doc.add_heading("5. Success criteria", level=1)
    _bullets(doc, [_txt(s) for s in (syn.get("success_criteria") or [])] or ["(none recorded)"])

    doc.add_heading("6. Out of scope", level=1)
    _bullets(doc, [_txt(s) for s in (syn.get("non_goals") or [])] or ["(none recorded)"])

    kt = syn.get("kill_tests") or {}
    if isinstance(kt, dict) and kt:
        doc.add_heading("7. Kill tests", level=1)
        _kv_table(doc, [(k.replace("_", " ").capitalize(),
                         f"{(v or {}).get('result', '')}: {(v or {}).get('note', '')}".strip(": "))
                        for k, v in kt.items() if isinstance(v, dict)])

    items = ps.get("open_items") or []
    doc.add_heading("8. Open items", level=1)
    if not items:
        _para(doc, "None.", after=8)
    for oi in items:
        res = oi.get("resolution") or {}
        p = _para(doc, after=2)
        r = p.add_run(f"{oi.get('item_id', '')}  "); r.bold = True; r.font.color.rgb = ACC
        _rich(p, _txt(oi.get("description")))
        state = (f"Resolved by {res.get('by', '')}: {res.get('rationale', '')}" if res
                 else ("Blocking" if oi.get("blocking") else "Open, not blocking"))
        _para(doc, state, size=9, color=DIM, indent_cm=0.6, after=6)

    if amends:
        doc.add_heading("9. Scope amendments", level=1)
        for a in amends:
            p = _para(doc, after=2)
            r = p.add_run(f"{a.get('amendment_id', '')} · v{a.get('from_version', '')} → v{a.get('to_version', '')} · "
                          f"{_date(a.get('timestamp'))} · {_who(a.get('amended_by'))} · {a.get('decision_log_ref', '')}")
            r.bold = True; r.font.size = Pt(9)
            _para(doc, _txt(a.get("rationale")), size=9, indent_cm=0.6, after=2)
            if a.get("previous_text"):
                _para(doc, "Previous scope: " + a["previous_text"], size=8.5, color=DIM, italic=True,
                      indent_cm=0.6, after=6)

    doc.add_heading(("10." if amends else "9.") + " Ratification", level=1)
    if ratified:
        _para(doc, f"Ratified by {_who(rat.get('ratified_by'))} on {_date(rat.get('timestamp'))} "
                   f"({rat.get('decision_log_ref', '')}).", after=2)
        if _txt(rat.get("rationale")):
            _para(doc, "Rationale: " + _txt(rat.get("rationale")), size=9, color=DIM)
    else:
        _para(doc, "Not yet ratified. Blocking open items must be resolved before the Program Owner can ratify.")

    _para(doc, ("Advisory working paper produced with AI assistance under human governance. Not operative law "
                "or a statement of any authority's policy."), size=8, color=DIM, italic=True, after=0)

    if not share:
        h = doc.add_heading("Appendix A. Interview transcript", level=1)
        h.paragraph_format.page_break_before = True
        consent = ((ps.get("interview") or {}).get("consent") or {})
        _para(doc, (f"Verbatim record of the purpose interview, {len(turns)} turns, from the restricted store "
                    f"({(ps.get('interview') or {}).get('transcript_ref') or 'restricted/interview.json'}). "
                    f"Transcript notice given: {'yes' if consent.get('notice_given') else 'not recorded'}. "
                    f"Excerpts consented for publication: {consent.get('publication_excerpts') or 'not recorded'}."),
              size=9, color=DIM, after=10)
        if not turns:
            _para(doc, "No transcript is on file for this program.")
        for i, m in enumerate(turns, 1):
            who = "Respondent" if m.get("role") == "user" else "Interviewer"
            hp = _para(doc, after=1)
            hp.paragraph_format.keep_with_next = True
            r = hp.add_run(f"{i}. {who}"); r.bold = True; r.font.size = Pt(8.5)
            r.font.color.rgb = INK if who == "Respondent" else DIM
            body = str(m.get("content", "")).replace(COMPLETE_MARKER, "").strip()
            lines = [ln for ln in body.splitlines()]
            bp = _para(doc, indent_cm=0.4, after=7)
            first = True
            for ln in lines:
                if not ln.strip():
                    continue
                if not first:
                    bp.add_run().add_break()
                _rich(bp, ln.strip(), size=9.5, color=INK if who == "Respondent" else None)
                first = False

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

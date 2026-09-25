"""Corpus Manifest export (.docx) — the source corpus as a readable, citable record.

For each source: its identity in the manifest (title, issuer, family, status,
evidence role, citation/locator, URL), how its text was obtained (fetched from
which URL and when, or uploaded as which document, with SHA-256 and sizes), any
URL correction or set-aside decision, and what distillation found in it.
Derived artifact only: governed/manifest/ and governed/corpus_texts/ are the
source of truth, and this document can be regenerated at any time.
"""
from __future__ import annotations

import io
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .defect_report import BAD, DIM, OK, _load, _para, _shade, styled_document

FAMILY_ORDER = ["statute", "regulation", "reporting_instruction", "guidance"]
ROLE_LABELS = {"primary_authority": "Primary authority", "binding_authority": "Binding authority",
               "authoritative_interpretation": "Authoritative interpretation",
               "nonbinding_guidance": "Non-binding guidance", "reference": "Reference",
               "cross_cutting_constraint": "Cross-cutting constraint",
               "enforcement_evidence": "Enforcement evidence", "reform_proposal": "Reform proposal"}


class NoManifest(Exception):
    pass


def _label(s: str | None) -> str:
    s = (s or "").replace("_", " ").strip()
    return s[:1].upper() + s[1:]


def _date(ts: str | None) -> str:
    if not ts:
        return ""
    try:
        return datetime.fromisoformat(ts.replace("Z", "+00:00")).strftime("%-d %b %Y %H:%M UTC")
    except ValueError:
        return ts[:16]


def _kb(n) -> str:
    return f"{int(n) / 1024:,.0f} KB" if n else ""


def _field_table(doc, rows: list[tuple[str, str, object]]) -> None:
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"; t.autofit = False
    for label, value, color in rows:
        c0, c1 = t.add_row().cells
        c0.width, c1.width = Cm(3.8), Cm(12.2)
        c0.text = label; _shade(c0, "F3F4F6")
        c1.text = value
        for c in (c0, c1):
            for r in c.paragraphs[0].runs:
                r.font.size = Pt(8.5)
        c0.paragraphs[0].runs[0].bold = True
        if color is not None:
            for r in c1.paragraphs[0].runs:
                r.font.color.rgb = color
    t.columns[0].width, t.columns[1].width = Cm(3.8), Cm(12.2)   # grid widths (LibreOffice/Google honour these)
    for row in t.rows[:-1]:                                       # keep each source's table on one page
        for c in row.cells:
            for par in c.paragraphs:
                par.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_docx(program_dir: str | Path, program_id: str) -> bytes:
    g = Path(program_dir) / "governed"
    man = _load(g / "manifest" / "manifest.json")
    if not man or not man.get("items"):
        raise NoManifest("No corpus manifest yet — add sources on the Corpus tab first")
    items = man["items"]
    acq = ((_load(g / "corpus_texts" / "acquisition.json") or {}).get("items")) or {}
    overrides = _load(g / "corpus_texts" / "url_overrides.json") or {}
    excluded = ((_load(g / "excluded_sources.json") or {}).get("items")) or {}
    ext = ((_load(g / "blueprint" / "extraction_register.json") or {}).get("items")) or {}
    ps = _load(g / "purpose_statement.json") or {}
    scope = (((ps.get("synthesis") or {}).get("scope_sentence") or {}).get("text") or "").strip()
    freeze = None
    log = g / "decisions.log.jsonl"
    if log.exists():
        for line in log.read_text().splitlines():
            if line.strip():
                e = json.loads(line)
                if e.get("type") == "manifest_freeze":
                    freeze = e

    frozen = bool(man.get("frozen"))
    n_excl = sum(1 for i in items if i["item_id"] in excluded)
    status = Counter((acq.get(i["item_id"]) or {}).get("status", "pending") for i in items)
    n_fetched, n_manual = status.get("fetched", 0), status.get("manual", 0)
    n_err, n_pend = status.get("error", 0), status.get("pending", 0)
    today = datetime.now(timezone.utc).strftime("%-d %B %Y")

    doc = styled_document(f"{program_id} — Corpus Manifest · page ", program_dir)
    doc.add_heading(f"{program_id} — Corpus Manifest", level=0)
    _para(doc, f"Rulebook Workbench · source corpus (Phase 1) · generated {today}", size=9, color=DIM)
    if not frozen:
        _para(doc, "DRAFT — this corpus is not yet frozen; the list below can still change.", bold=True, color=BAD)
    _para(doc, ("This is the record of every source in the program's corpus: what it is, where it comes "
                "from, and exactly which text the Workbench read. Fetched sources carry the URL and the time "
                "they were retrieved; uploaded sources carry the document's file name and SHA-256 fingerprint, "
                "so any reader can confirm they hold the same text. Distillation reads only these texts."))
    if scope:
        p = _para(doc); p.add_run("Scope. ").bold = True; p.add_run(scope)

    doc.add_heading("Corpus at a glance", level=2)
    facts = [
        ("Status", (f"Frozen {_date(man.get('frozen_at'))}" + (
            f" by {freeze['decided_by']['name']} ({freeze['decided_by']['role']}), {freeze['entry_id']}"
            if freeze else "")) if frozen else "Draft (not frozen)", None),
    ]
    if freeze and freeze.get("rationale"):
        facts.append(("Freeze rationale", freeze["rationale"], None))
    facts += [
        ("Content hash", man.get("content_hash") or "— (set at freeze)", None),
        ("Sources", f"{len(items)} in the manifest · {len(items) - n_excl} in the working corpus · "
                    f"{n_excl} set aside", None),
        ("Texts", f"{n_fetched} fetched from the web · {n_manual} uploaded documents"
                  + (f" · {n_err} errors" if n_err else "") + (f" · {n_pend} pending" if n_pend else ""),
         BAD if (n_err or n_pend) else None),
    ]
    if overrides:
        facts.append(("URL corrections", f"{len(overrides)} source(s) fetched from a corrected URL (shown below)", None))
    _field_table(doc, facts)

    fams = sorted({i.get("family", "") for i in items},
                  key=lambda f: (FAMILY_ORDER.index(f) if f in FAMILY_ORDER else 99, f))
    t = doc.add_table(rows=1, cols=5); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [Cm(5.2), Cm(2.7), Cm(2.7), Cm(2.7), Cm(2.7)]
    for i, h in enumerate(["Family", "Sources", "Fetched", "Uploaded", "Set aside"]):
        c = t.rows[0].cells[i]; c.text = h; _shade(c, "F3F4F6"); c.paragraphs[0].runs[0].bold = True
    for fam in fams:
        its = [i for i in items if i.get("family") == fam]
        st = Counter((acq.get(i["item_id"]) or {}).get("status") for i in its)
        row = t.add_row().cells
        vals = [_label(fam), len(its), st.get("fetched", 0), st.get("manual", 0),
                sum(1 for i in its if i["item_id"] in excluded)]
        for k, v in enumerate(vals):
            row[k].text = str(v) if (k == 0 or v) else "–"
    for k, col in enumerate(t.columns):
        col.width = widths[k]
    for r_ in t.rows:
        for k, c in enumerate(r_.cells):
            c.width = widths[k]
            for p in c.paragraphs:
                if k: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs: run.font.size = Pt(9)
    _para(doc, ("Advisory working paper produced with AI assistance under human governance. Titles, issuers "
                "and labels are the program's own; check citations against the issuing authority before "
                "relying on them."), size=8, color=DIM, italic=True, after=0)

    for fam in fams:
        its = [i for i in items if i.get("family") == fam]
        h = doc.add_heading(f"{_label(fam)} ({len(its)} sources)", level=1)
        h.paragraph_format.page_break_before = True
        for item in its:
            iid = item["item_id"]
            a = acq.get(iid) or {}
            hp = doc.add_heading(item.get("title", iid), level=3)
            hp.paragraph_format.keep_with_next = True
            rows = [("Issuer", item.get("issuer", ""), None),
                    ("Status", _label(item.get("status")) + (
                        f" · evidence role: {ROLE_LABELS.get(item.get('evidence_role'), _label(item.get('evidence_role')))}"
                        if item.get("evidence_role") else ""), None),
                    ("Citation", item.get("locator", ""), None)]
            if item.get("url"):
                rows.append(("URL (manifest)", item["url"], None))
            if a.get("status") == "fetched":
                src = a.get("source_url", "")
                if iid in overrides:
                    rows.append(("URL corrected", f"Fetched from {src} — the manifest URL no longer resolved", None))
                rows.append(("Text obtained", f"Fetched {_date(a.get('fetched_at'))} from {src} · "
                             f"{_kb(a.get('raw_bytes'))} {a.get('raw_ext', '').upper()} → "
                             f"{a.get('text_chars', 0):,} characters of text", None))
                rows.append(("SHA-256 (as fetched)", a.get("raw_sha256", "").removeprefix("sha256:"), None))
            elif a.get("status") == "manual":
                rows.append(("Text obtained", f"Uploaded document “{a.get('filename', '')}”, "
                             f"{_date(a.get('uploaded_at'))} · {_kb(a.get('raw_bytes'))} → "
                             f"{a.get('text_chars', 0):,} characters of text", None))
                rows.append(("SHA-256 (as uploaded)", a.get("raw_sha256", "").removeprefix("sha256:"), None))
            elif a.get("status") == "error":
                rows.append(("Text obtained", "Not acquired — " + "; ".join(a.get("errors", []))[:300], BAD))
            else:
                rows.append(("Text obtained", "Not yet acquired", BAD))
            if iid in excluded:
                rows.append(("Set aside", excluded[iid].get("reason", ""), BAD))
            e = ext.get(iid)
            if e and e.get("status") == "extracted":
                rows.append(("Distilled", "Nothing within the program's scope" if e.get("nothing_in_scope")
                             else f"{e.get('obligations', 0)} obligations · {e.get('definitions', 0)} definitions"
                                  + (f" · {e.get('citations_verified', 0)}/{e.get('citations_total', 0)} quotes verified"
                                     if e.get("citations_total") else ""),
                             DIM if e.get("nothing_in_scope") else OK))
            rows.append(("Workbench ID", iid, DIM))
            _field_table(doc, rows)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

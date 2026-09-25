"""Defect Register export (.docx) — a readable, shareable rendering of the
Phase 2 Defect Register (registers/defects.json).

Derived artifact only: the JSON register is the source of truth, and this
document can be regenerated at any time. Finding references (run_id#index)
match the identifiers the refactor pass uses.
"""
from __future__ import annotations

import io
import json
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TAXONOMY = {
    "D1": "Conflicting requirements", "D2": "Divergent definitions", "D3": "Duplicate provisions",
    "D4": "Undefined material term", "D5": "Dead or dangling reference",
    "D6": "Superseded in substance but never revoked", "D7": "Scattered requirement",
    "D8": "Obsolete or archaic provision", "D9": "Gap", "D10": "Applicability inconsistency",
}
RUN_LABELS = {"defects-cross": "Cross-corpus"}
INK, DIM, ACC = RGBColor(0x1F, 0x29, 0x37), RGBColor(0x6B, 0x72, 0x80), RGBColor(0xB4, 0x53, 0x09)
OK, BAD, QUOTE = RGBColor(0x15, 0x80, 0x3D), RGBColor(0xB9, 0x1C, 0x1C), RGBColor(0x37, 0x41, 0x51)


class NoDefectRegister(Exception):
    pass


def _run_label(run_id: str) -> str:
    if run_id in RUN_LABELS:
        return RUN_LABELS[run_id]
    fam = run_id.removeprefix("defects-").replace("_", " ")
    return fam[:1].upper() + fam[1:]


def _load(path: Path) -> dict | None:
    return json.loads(path.read_text()) if path.exists() else None


def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _page_field(paragraph) -> None:
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)
    run.font.size = Pt(8); run.font.color.rgb = DIM


def _para(doc, text="", *, size=None, color=None, bold=False, italic=False, after=6, indent_cm=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if indent_cm:
        p.paragraph_format.left_indent = Cm(indent_cm)
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color
    return p


def styled_document(footer_text: str):
    """A4 Word document with the Workbench house style and a 'page N' footer."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.54))
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(10)
    st.font.color.rgb = INK
    for name, size in (("Normal", 10), ("Title", 20), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5)):
        sty = doc.styles[name]; sty.font.name, sty.font.size = "Arial", Pt(size)
        if name != "Normal":
            sty.font.bold = True
        sty.font.color.rgb = INK
        rfonts = sty.element.get_or_add_rPr().get_or_add_rFonts()   # beat the theme (Calibri) font
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            rfonts.attrib.pop(qn(attr), None)                          # theme fonts win over names
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), "Arial")

    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run(footer_text)
    fr.font.size = Pt(8); fr.font.color.rgb = DIM
    _page_field(fp)
    return doc


def build_docx(program_dir: str | Path, program_id: str) -> bytes:
    g = Path(program_dir) / "governed"
    reg = _load(g / "registers" / "defects.json")
    runs = (reg or {}).get("runs") or {}
    if not runs:
        raise NoDefectRegister("No Defect Register yet — run defect detection on the Derived Blueprint first")
    man = _load(g / "manifest" / "manifest.json") or {}
    titles = {i["item_id"]: i.get("title", i["item_id"]) for i in man.get("items", [])}
    ps = _load(g / "purpose_statement.json") or {}
    scope = (((ps.get("synthesis") or {}).get("scope_sentence") or {}).get("text") or "").strip()
    excluded = set(((_load(g / "excluded_sources.json") or {}).get("items") or {}).keys())
    n_sources = len([i for i in man.get("items", []) if i["item_id"] not in excluded])

    run_ids = sorted(runs)
    total = sum(len(runs[r].get("findings", [])) for r in run_ids)
    unverified = sum(1 for r in run_ids for f in runs[r].get("findings", [])
                     if any(not l.get("verified") for l in f.get("locations", [])))
    today = datetime.now(timezone.utc).strftime("%-d %B %Y")

    doc = styled_document(f"{program_id} — Defect Register · advisory working paper · page ")

    doc.add_heading(f"{program_id} — Defect Register", level=0)
    _para(doc, f"Rulebook Workbench · Derived Blueprint (Phase 2) · generated {today}", size=9, color=DIM)
    _para(doc, (f"This register lists the {total} defects the Workbench found while distilling the {n_sources} "
                "in-scope sources of this programme. Each finding names a defect type, explains the problem, "
                "and cites the sources where it appears, with the verbatim passage. A tick means the quoted "
                "passage was checked by code and found word-for-word in the source text."))
    _para(doc, ("Defects are observations about the current rules, not errors in the analysis. In the Refactor "
                "pass each one is worked through: a fix is proposed and a named reviewer accepts, amends or "
                "rejects it. References such as defects-cross#0 match the Workbench's own identifiers."))
    if scope:
        p = _para(doc); p.add_run("Scope. ").bold = True; p.add_run(scope)

    doc.add_heading("Findings by type and detection pass", level=2)
    t = doc.add_table(rows=1, cols=len(run_ids) + 2)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = ["Defect type"] + [_run_label(r) for r in run_ids] + ["Total"]
    for i, h in enumerate(hdr):
        c = t.rows[0].cells[i]; c.text = h; _shade(c, "F3F4F6")
        c.paragraphs[0].runs[0].bold = True
    for code, name in TAXONOMY.items():
        per = [sum(1 for f in runs[r].get("findings", []) if f.get("code") == code) for r in run_ids]
        row = t.add_row().cells
        row[0].text = f"{code} {name}"
        for i, n in enumerate(per):
            row[i + 1].text = str(n) if n else "–"
        row[-1].text = str(sum(per)) if sum(per) else "–"
    row = t.add_row().cells
    row[0].text = "All findings"
    for i, r in enumerate(run_ids):
        row[i + 1].text = str(len(runs[r].get("findings", [])))
    row[-1].text = str(total)
    for c in row:
        _shade(c, "F3F4F6"); c.paragraphs[0].runs[0].bold = True
    t.autofit = False
    first, rest = Cm(5.2), Cm((16.0 - 5.2) / (len(run_ids) + 1))
    for col_i, col in enumerate(t.columns):
        col.width = first if col_i == 0 else rest
    for r_ in t.rows:
        for i, c in enumerate(r_.cells):
            c.width = first if i == 0 else rest
            for p in c.paragraphs:
                if i: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs: run.font.size = Pt(9)
    _para(doc, (f"{unverified} of the {total} findings include at least one quoted passage that was not matched "
                "word-for-word in the source; these are marked in red and should be checked before any fix is "
                "accepted. The same issue can surface in both a family pass and the cross-corpus pass."),
          size=9, color=DIM, after=8)
    doc.add_heading("How this register is organised", level=2)
    _para(doc, "Findings are grouped by the detection pass that found them, each starting on a new page:", after=3)
    for r in run_ids:
        _para(doc, f"{_run_label(r)} — {len(runs[r].get('findings', []))} findings", indent_cm=0.6, after=2)
    _para(doc, ("Advisory working paper produced with AI assistance under human governance. Not operative law "
                "or a statement of any authority's policy."), size=8, color=DIM, italic=True, after=0)

    for r in run_ids:
        run = runs[r]
        h = doc.add_heading(f"{_run_label(r)} ({len(run.get('findings', []))} findings)", level=1)
        h.paragraph_format.page_break_before = True
        _para(doc, f"Detection pass: {run.get('scope_label') or r} · run "
                   f"{str(run.get('detected_at', ''))[:16].replace('T', ' ')} UTC", size=9, color=DIM)
        for i, f in enumerate(run.get("findings", [])):
            code = f.get("code", "?")
            hp = doc.add_heading(level=3)
            hp.paragraph_format.keep_with_next = True
            cr = hp.add_run(f"{code}  "); cr.font.color.rgb = ACC
            hp.add_run(f.get("title", ""))
            meta = _para(doc, f"{r}#{i} · {TAXONOMY.get(code, code)}", size=8, color=DIM, after=3)
            meta.paragraph_format.keep_with_next = True
            _para(doc, f.get("description", ""), after=5)
            for loc in f.get("locations", []):
                lp = _para(doc, indent_cm=0.6, after=2)
                lp.paragraph_format.keep_with_next = True
                a = lp.add_run("↳ "); a.font.color.rgb = DIM
                s = lp.add_run(titles.get(loc.get("item_id"), loc.get("item_id", ""))); s.bold = True
                s.font.size = Pt(9)
                v = lp.add_run("   ✓ quote verified in source" if loc.get("verified")
                               else "   quote not matched verbatim in source")
                v.font.size = Pt(8); v.font.color.rgb = OK if loc.get("verified") else BAD
                lines = [x for x in str(loc.get("quote", "")).splitlines() if x.strip()]
                if lines:
                    qp = _para(doc, indent_cm=1.2, after=6)
                    for k, line in enumerate(lines):
                        q = qp.add_run(("“" if k == 0 else "") + line + ("”" if k == len(lines) - 1 else ""))
                        q.italic = True; q.font.size = Pt(9); q.font.color.rgb = QUOTE
                        if k < len(lines) - 1:
                            q.add_break(WD_BREAK.LINE)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
